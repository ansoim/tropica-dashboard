"""Fetch the public Tropica workbook and refresh data/sales.json."""

from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from urllib.request import Request, urlopen
from io import BytesIO

from openpyxl import load_workbook


SHARE_URL = (
    "https://ansoim-my.sharepoint.com/:x:/p/abhishek_kumar/"
    "IQDq2wrkfrleRa0HdJ3RCWAUAS8UuAWbIxp4n8CXHA6J4o4?e=E83P9V"
)
ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "data" / "sales.json"


def download_url(share_url: str) -> str:
    """Ask OneDrive/SharePoint to return the workbook file, not its viewer page."""
    parts = urlsplit(share_url)
    query = dict(parse_qsl(parts.query, keep_blank_values=True))
    query["download"] = "1"
    return urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(query), ""))


def main() -> None:
    request = Request(
        download_url(SHARE_URL),
        headers={
            "User-Agent": "Mozilla/5.0 TropicaDashboard/1.0",
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
        },
    )
    with urlopen(request, timeout=45) as response:
        workbook_bytes = response.read()

    # XLSX files are ZIP containers. Fail clearly if SharePoint returned an HTML page.
    if not workbook_bytes.startswith(b"PK"):
        preview = workbook_bytes[:160].decode("utf-8", errors="replace").replace("\n", " ")
        raise RuntimeError(
            "SharePoint did not return an .xlsx download. Check that the link is set to "
            f"Anyone with the link can view and download. Response preview: {preview}"
        )

    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    if "Summary" not in workbook.sheetnames:
        raise RuntimeError("The workbook does not contain a sheet named Summary.")
    values = [list(row) for row in workbook["Summary"].iter_rows(
        min_row=4, max_row=10, min_col=1, max_col=6, values_only=True
    )]
    workbook.close()

    old_feed = {}
    if DATA_FILE.exists():
        try:
            old_feed = json.loads(DATA_FILE.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            old_feed = {}

    if old_feed.get("values") == values:
        print("Workbook values are unchanged; keeping the existing feed.")
        changed = False
    else:
        feed = {
            "updatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
            "values": values,
        }
        DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
        DATA_FILE.write_text(json.dumps(feed, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"Refreshed dashboard data from Summary!A4:F10 ({len(values)} rows).")
        changed = True

    output_path = os.environ.get("GITHUB_OUTPUT")
    if output_path:
        with open(output_path, "a", encoding="utf-8") as output:
            output.write(f"changed={'true' if changed else 'false'}\n")


if __name__ == "__main__":
    main()
